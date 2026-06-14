from io import BytesIO
from typing import Iterable, Mapping, Sequence

from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


def _fill_sheet(
    worksheet,
    headers: Sequence[str],
    rows: Iterable[Sequence[object]],
    number_formats: Mapping[str, str] | None = None,
) -> None:
    worksheet.append(list(headers))
    header_fill = PatternFill("solid", fgColor="EAF2FF")
    for cell in worksheet[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill

    for row in rows:
        worksheet.append(list(row))

    if number_formats:
        for column_index, header in enumerate(headers, start=1):
            number_format = number_formats.get(header)
            if not number_format:
                continue
            for row_index in range(2, worksheet.max_row + 1):
                worksheet.cell(row=row_index, column=column_index).number_format = number_format

    for column_cells in worksheet.columns:
        max_length = max(len(str(cell.value or "")) for cell in column_cells)
        worksheet.column_dimensions[get_column_letter(column_cells[0].column)].width = min(max(max_length + 2, 12), 32)


def build_xlsx_response(filename: str, headers: Sequence[str], rows: Iterable[Sequence[object]]) -> StreamingResponse:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Report"
    _fill_sheet(worksheet, headers, rows)

    stream = BytesIO()
    workbook.save(stream)
    stream.seek(0)

    return StreamingResponse(
        stream,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


def build_multi_sheet_xlsx_response(filename: str, sheets: Sequence[Mapping[str, object]]) -> StreamingResponse:
    workbook = Workbook()
    workbook.remove(workbook.active)
    for sheet in sheets:
        title = str(sheet["title"])[:31]
        worksheet = workbook.create_sheet(title=title)
        _fill_sheet(
            worksheet,
            sheet["headers"],  # type: ignore[arg-type]
            sheet["rows"],  # type: ignore[arg-type]
            sheet.get("number_formats"),  # type: ignore[arg-type]
        )

    stream = BytesIO()
    workbook.save(stream)
    stream.seek(0)

    return StreamingResponse(
        stream,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )
